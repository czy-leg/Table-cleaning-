"""表格清洗助手的数据读取、清洗和导出核心逻辑。"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from typing import Any, Hashable, Sequence

import pandas as pd
from openpyxl import load_workbook


CSV_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")
CSV_DELIMITERS = {
    ",": "逗号（,）",
    "\t": "制表符（Tab）",
    ";": "分号（;）",
    "|": "竖线（|）",
}


@dataclass(frozen=True)
class CleaningOptions:
    """一次清洗任务的可配置选项。"""

    drop_duplicate_rows: bool = False
    drop_blank_rows: bool = False
    drop_blank_columns: bool = False
    trim_whitespace: bool = False
    deduplicate_by_columns: bool = False
    duplicate_columns: tuple[Hashable, ...] = ()
    case_sensitive: bool = True

    def signature(self) -> tuple[Any, ...]:
        """返回可用于判断选项是否发生变化的稳定标识。"""

        return (
            self.drop_duplicate_rows,
            self.drop_blank_rows,
            self.drop_blank_columns,
            self.trim_whitespace,
            self.deduplicate_by_columns,
            self.duplicate_columns,
            self.case_sensitive,
        )


@dataclass(frozen=True)
class CleaningStats:
    """清洗前后的行列数与各项删除统计。"""

    original_rows: int
    original_columns: int
    result_rows: int
    result_columns: int
    removed_duplicate_rows: int
    removed_blank_rows: int
    removed_blank_columns: int


@dataclass(frozen=True)
class DecodedCsv:
    """成功解码后的 CSV 内容和使用的编码。"""

    text: str
    encoding: str


class CsvDecodeError(ValueError):
    """CSV 无法使用常见编码解码。"""


def _is_blank_cell(value: Any) -> bool:
    """判断单元格是否为空、缺失或只含空白字符。"""

    if value is None:
        return True

    if isinstance(value, str):
        return value.strip() == ""

    try:
        missing = pd.isna(value)
    except (TypeError, ValueError):
        return False

    return bool(missing) if isinstance(missing, bool) else False


def trim_text_cells(dataframe: pd.DataFrame) -> pd.DataFrame:
    """只清理文本单元格首尾空白，保留非文本值及其类型。"""

    result = dataframe.copy(deep=True)
    for column in result.columns:
        series = result[column]
        result[column] = series.map(
            lambda value: value.strip() if isinstance(value, str) else value
        )
    return result


def drop_blank_rows(dataframe: pd.DataFrame) -> pd.DataFrame:
    """删除所有单元格都为空的整行。"""

    if dataframe.empty:
        return dataframe.copy(deep=True)

    blank_mask = dataframe.apply(
        lambda row: all(_is_blank_cell(value) for value in row),
        axis=1,
    )
    return dataframe.loc[~blank_mask].copy(deep=True)


def drop_blank_columns(dataframe: pd.DataFrame) -> pd.DataFrame:
    """删除所有单元格都为空的整列。"""

    if dataframe.empty:
        return dataframe.copy(deep=True)

    blank_columns = [
        column
        for column in dataframe.columns
        if all(_is_blank_cell(value) for value in dataframe[column].tolist())
    ]
    return dataframe.drop(columns=blank_columns).copy(deep=True)


def drop_duplicate_rows(
    dataframe: pd.DataFrame,
    *,
    subset: Sequence[Hashable] | None = None,
    case_sensitive: bool = True,
) -> pd.DataFrame:
    """删除重复行，可指定列并选择是否区分文本大小写。"""

    if dataframe.empty:
        return dataframe.copy(deep=True)

    key_columns = list(subset) if subset is not None else None
    if key_columns is not None:
        if not key_columns:
            raise ValueError("请至少选择一列用于重复判断。")
        missing_columns = [
            column for column in key_columns if column not in dataframe.columns
        ]
        if missing_columns:
            missing_text = "、".join(str(column) for column in missing_columns)
            raise ValueError(f"找不到用于重复判断的列：{missing_text}")

    comparison_data = dataframe
    if not case_sensitive:
        comparison_data = dataframe.copy(deep=True)
        target_columns = key_columns if key_columns is not None else list(comparison_data.columns)
        for column in target_columns:
            comparison_data[column] = comparison_data[column].map(
                lambda value: value.casefold() if isinstance(value, str) else value
            )

    duplicate_mask = comparison_data.duplicated(
        subset=key_columns,
        keep="first",
    )
    return dataframe.loc[~duplicate_mask].copy(deep=True)


def clean_dataframe(
    dataframe: pd.DataFrame,
    options: CleaningOptions,
) -> tuple[pd.DataFrame, CleaningStats]:
    """按选项执行清洗，返回新数据和完整统计。"""

    if not isinstance(dataframe, pd.DataFrame):
        raise TypeError("待清洗数据必须是 pandas DataFrame。")

    original_rows = len(dataframe)
    original_columns = len(dataframe.columns)
    result = dataframe.copy(deep=True)
    removed_duplicate_rows = 0
    removed_blank_rows = 0
    removed_blank_columns = 0

    if options.trim_whitespace:
        result = trim_text_cells(result)

    if options.drop_blank_columns:
        before_columns = len(result.columns)
        result = drop_blank_columns(result)
        removed_blank_columns += before_columns - len(result.columns)

    if options.drop_blank_rows:
        before_rows = len(result)
        result = drop_blank_rows(result)
        removed_blank_rows += before_rows - len(result)

    if options.drop_duplicate_rows:
        before_rows = len(result)
        result = drop_duplicate_rows(
            result,
            case_sensitive=options.case_sensitive,
        )
        removed_duplicate_rows += before_rows - len(result)

    if options.deduplicate_by_columns:
        if not options.duplicate_columns:
            raise ValueError("已启用“按勾选列删除重复”，请至少选择一列。")
        before_rows = len(result)
        result = drop_duplicate_rows(
            result,
            subset=options.duplicate_columns,
            case_sensitive=options.case_sensitive,
        )
        removed_duplicate_rows += before_rows - len(result)

    stats = CleaningStats(
        original_rows=original_rows,
        original_columns=original_columns,
        result_rows=len(result),
        result_columns=len(result.columns),
        removed_duplicate_rows=removed_duplicate_rows,
        removed_blank_rows=removed_blank_rows,
        removed_blank_columns=removed_blank_columns,
    )
    return result, stats


def decode_csv_bytes(data: bytes) -> DecodedCsv:
    """依次尝试 UTF-8 BOM、UTF-8 和 GB18030 解码 CSV。"""

    if not data:
        raise CsvDecodeError("CSV 文件为空。")

    failures: list[str] = []
    for encoding in CSV_ENCODINGS:
        try:
            return DecodedCsv(data.decode(encoding), encoding)
        except UnicodeDecodeError as exc:
            failures.append(f"{encoding}（第 {exc.start + 1} 字节）")

    attempted = "、".join(failures)
    raise CsvDecodeError(
        f"无法识别 CSV 编码，已尝试：{attempted}。"
        "请用记事本或表格软件将文件另存为 UTF-8 或 GB18030 后重试。"
    )


def detect_csv_delimiter(text: str) -> str | None:
    """尝试从 CSV 文本前部识别逗号、制表符、分号或竖线分隔符。"""

    sample = text[:16_384]
    if not sample.strip():
        return None

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="".join(CSV_DELIMITERS))
    except csv.Error:
        return None
    return dialect.delimiter if dialect.delimiter in CSV_DELIMITERS else None


def read_csv_bytes(
    data: bytes,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
) -> pd.DataFrame:
    """读取内存中的 CSV 数据，保留空白行并让 pandas 谨慎推断数据类型。"""

    decoded = decode_csv_bytes(data) if encoding is None else None
    text = decoded.text if decoded is not None else data.decode(encoding)
    delimiter = delimiter or detect_csv_delimiter(text)
    if delimiter is None:
        raise ValueError("未能自动识别 CSV 分隔符，请手动选择分隔符。")
    if delimiter not in CSV_DELIMITERS:
        raise ValueError("不支持该 CSV 分隔符。")
    if not text.strip():
        raise ValueError("CSV 文件没有可读取的内容。")

    try:
        dataframe = pd.read_csv(
            StringIO(text),
            sep=delimiter,
            keep_default_na=False,
            na_filter=False,
            skip_blank_lines=False,
        )
    except (pd.errors.ParserError, UnicodeDecodeError, ValueError) as exc:
        raise ValueError(
            "CSV 内容无法按所选编码和分隔符读取，请检查文件格式。"
        ) from exc

    if dataframe.empty and len(dataframe.columns) == 0:
        raise ValueError("CSV 文件没有可读取的表格内容。")
    return dataframe


def list_excel_sheets(data: bytes) -> list[str]:
    """读取 Excel 工作簿的工作表名称，不执行宏。"""

    try:
        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            keep_vba=False,
            data_only=False,
        )
    except Exception as exc:
        raise ValueError(
            "无法读取 Excel 文件。请确认文件未加密、未损坏，且不是仅改扩展名的其他格式。"
        ) from exc

    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_excel_bytes(data: bytes, *, sheet_name: str) -> pd.DataFrame:
    """读取指定 Excel 工作表；宏不会执行，公式以公式文本读取。"""

    try:
        dataframe = pd.read_excel(
            BytesIO(data),
            sheet_name=sheet_name,
            engine="openpyxl",
            engine_kwargs={
                "data_only": False,
                "keep_vba": False,
            },
        )
    except Exception as exc:
        raise ValueError(
            f"无法读取工作表“{sheet_name}”。请确认文件未加密、未损坏，且工作表存在。"
        ) from exc
    return dataframe


_FORMULA_PREFIX_PATTERN = re.compile(r"^[=+\-@\t\r]")


def _looks_like_formula_text(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    return bool(_FORMULA_PREFIX_PATTERN.match(value.lstrip(" \t\r\n")))


def protect_formula_values(dataframe: pd.DataFrame) -> pd.DataFrame:
    """在危险文本前添加单引号，避免被 Excel 当作公式执行。"""

    result = dataframe.copy(deep=True)
    for column in result.columns:
        result[column] = result[column].map(
            lambda value: f"'{value}" if _looks_like_formula_text(value) else value
        )
    result.columns = [
        f"'{column}" if _looks_like_formula_text(column) else column
        for column in result.columns
    ]
    return result


def export_csv_bytes(
    dataframe: pd.DataFrame,
    *,
    utf8_bom: bool = True,
    protect_formulas: bool = False,
) -> bytes:
    """把 DataFrame 导出为 CSV 字节内容。"""

    output_data = (
        protect_formula_values(dataframe) if protect_formulas else dataframe
    )
    text = output_data.to_csv(index=False, lineterminator="\n")
    return text.encode("utf-8-sig" if utf8_bom else "utf-8")


def export_excel_bytes(
    dataframe: pd.DataFrame,
    *,
    protect_formulas: bool = False,
) -> bytes:
    """把 DataFrame 导出为不含宏的 XLSX 文件。"""

    output_data = (
        protect_formula_values(dataframe) if protect_formulas else dataframe
    )
    buffer = BytesIO()
    try:
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            output_data.to_excel(writer, index=False, sheet_name="清洗结果")
    except Exception as exc:
        raise ValueError(
            "Excel 导出失败。请检查数据中是否包含无法写入工作簿的内容。"
        ) from exc
    return buffer.getvalue()


def output_filename(original_name: str, extension: str) -> str:
    """根据原文件名生成安全的 _cleaned 下载文件名。"""

    from pathlib import Path

    stem = Path(original_name).stem.strip() or "data"
    extension = extension.lower().lstrip(".")
    if extension not in {"csv", "xlsx"}:
        raise ValueError("仅支持导出 CSV 或 XLSX。")
    return f"{stem}_cleaned.{extension}"
