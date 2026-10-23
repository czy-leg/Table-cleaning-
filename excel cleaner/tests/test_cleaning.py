"""表格清洗助手核心逻辑测试。"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from pandas.testing import assert_frame_equal
import pytest

from cleaner import (
    CleaningOptions,
    clean_dataframe,
    decode_csv_bytes,
    drop_blank_columns,
    drop_blank_rows,
    drop_duplicate_rows,
    export_csv_bytes,
    export_excel_bytes,
    list_excel_sheets,
    protect_formula_values,
    read_csv_bytes,
    read_excel_bytes,
    trim_text_cells,
)


def test_drop_duplicate_rows() -> None:
    original = pd.DataFrame(
        {
            "姓名": ["张三", "张三", "李四"],
            "城市": ["北京", "北京", "上海"],
        }
    )

    result = drop_duplicate_rows(original)

    assert len(result) == 2
    assert result["姓名"].tolist() == ["张三", "李四"]


def test_drop_blank_rows_handles_missing_empty_and_spaces() -> None:
    original = pd.DataFrame(
        {
            "姓名": ["张三", None, "", "  "],
            "城市": ["北京", None, "   ", None],
        }
    )

    result = drop_blank_rows(original)

    assert len(result) == 1
    assert result.iloc[0].tolist() == ["张三", "北京"]


def test_drop_blank_columns() -> None:
    original = pd.DataFrame(
        {
            "姓名": ["张三", "李四"],
            "空白列一": [None, ""],
            "空白列二": ["  ", "   "],
        }
    )

    result = drop_blank_columns(original)

    assert result.columns.tolist() == ["姓名"]


def test_trim_text_does_not_change_non_text_values() -> None:
    original = pd.DataFrame(
        {
            "文本": ["  张三  ", " 北京 "],
            "数字": [1, 2.5],
            "日期": [date(2026, 1, 1), date(2026, 1, 2)],
        }
    )

    result = trim_text_cells(original)

    assert result["文本"].tolist() == ["张三", "北京"]
    assert result["数字"].tolist() == [1, 2.5]
    assert result["日期"].tolist() == original["日期"].tolist()
    assert pd.api.types.is_numeric_dtype(result["数字"])


def test_deduplicate_by_selected_columns_and_ignore_case() -> None:
    original = pd.DataFrame(
        {
            "编号": ["A-01", "a-01", "A-02"],
            "说明": ["第一次", "第二次", "第三次"],
        }
    )

    result = drop_duplicate_rows(
        original,
        subset=["编号"],
        case_sensitive=False,
    )

    assert len(result) == 2
    assert result["说明"].tolist() == ["第一次", "第三次"]


def test_empty_dataframe_is_supported() -> None:
    options = CleaningOptions(
        drop_duplicate_rows=True,
        drop_blank_rows=True,
        drop_blank_columns=True,
        trim_whitespace=True,
    )

    result, stats = clean_dataframe(pd.DataFrame(), options)

    assert result.empty
    assert stats.original_rows == 0
    assert stats.original_columns == 0
    assert stats.removed_duplicate_rows == 0
    assert stats.removed_blank_rows == 0
    assert stats.removed_blank_columns == 0


def test_chinese_csv_reads_utf8_bom() -> None:
    content = "姓名,城市\n张三,北京\n李四,上海\n".encode("utf-8-sig")

    decoded = decode_csv_bytes(content)
    dataframe = read_csv_bytes(content, encoding=decoded.encoding)

    assert decoded.encoding == "utf-8-sig"
    assert dataframe.columns.tolist() == ["姓名", "城市"]
    assert dataframe.iloc[1].tolist() == ["李四", "上海"]


def test_chinese_csv_reads_gb18030() -> None:
    content = "姓名,城市\n王五,广州\n".encode("gb18030")

    decoded = decode_csv_bytes(content)
    dataframe = read_csv_bytes(content, encoding=decoded.encoding)

    assert dataframe.iloc[0].tolist() == ["王五", "广州"]


def test_cleaning_statistics_are_reported() -> None:
    original = pd.DataFrame(
        {
            "姓名": ["张三", "张三", None, "李四"],
            "城市": ["北京", "北京", None, "上海"],
            "空列": [None, None, None, None],
        }
    )
    options = CleaningOptions(
        drop_duplicate_rows=True,
        drop_blank_rows=True,
        drop_blank_columns=True,
    )

    result, stats = clean_dataframe(original, options)

    assert result.shape == (2, 2)
    assert stats.original_rows == 4
    assert stats.original_columns == 3
    assert stats.result_rows == 2
    assert stats.result_columns == 2
    assert stats.removed_duplicate_rows == 1
    assert stats.removed_blank_rows == 1
    assert stats.removed_blank_columns == 1


def test_original_dataframe_is_not_mutated() -> None:
    original = pd.DataFrame(
        {
            "姓名": ["  张三  ", "张三", None],
            "城市": ["北京", "北京", None],
        }
    )
    snapshot = original.copy(deep=True)
    options = CleaningOptions(
        trim_whitespace=True,
        drop_duplicate_rows=True,
        drop_blank_rows=True,
    )

    clean_dataframe(original, options)

    assert_frame_equal(original, snapshot)


def test_csv_export_round_trip() -> None:
    original = pd.DataFrame(
        {
            "姓名": ["张三", "李四"],
            "金额": [12.5, 30],
            "城市": ["北京", "上海"],
        }
    )

    exported = export_csv_bytes(original, utf8_bom=True)
    reread = read_csv_bytes(exported, encoding="utf-8-sig", delimiter=",")

    assert exported.startswith(b"\xef\xbb\xbf")
    assert_frame_equal(reread, original)


def test_formula_protection_changes_only_risky_text() -> None:
    original = pd.DataFrame(
        {
            "文本": ["=1+1", "正常文本", "-文字", -10, "@命令"],
            "数字": [1, 2, 3, 4, 5],
        }
    )

    result = protect_formula_values(original)

    assert result["文本"].tolist() == ["'=1+1", "正常文本", "'-文字", -10, "'@命令"]
    assert result["数字"].tolist() == original["数字"].tolist()


def test_formula_protection_includes_column_headers() -> None:
    original = pd.DataFrame({"=危险标题": ["正常内容"]})

    result = protect_formula_values(original)

    assert result.columns.tolist() == ["'=危险标题"]


def test_excel_export_creates_xlsx_file() -> None:
    original = pd.DataFrame({"姓名": ["张三"], "金额": [10]})

    exported = export_excel_bytes(original)
    workbook = load_workbook(BytesIO(exported), read_only=True)

    try:
        assert workbook.sheetnames == ["清洗结果"]
        worksheet = workbook["清洗结果"]
        assert list(worksheet.values) == [
            ("姓名", "金额"),
            ("张三", 10),
        ]
    finally:
        workbook.close()


def test_excel_sheet_reading_keeps_formula_text() -> None:
    workbook = Workbook()
    first_sheet = workbook.active
    first_sheet.title = "说明"
    first_sheet["A1"] = "示例"

    data_sheet = workbook.create_sheet("数据")
    data_sheet.append(["名称", "金额"])
    data_sheet.append(["合计", "=SUM(B3:B4)"])
    buffer = BytesIO()
    workbook.save(buffer)

    sheets = list_excel_sheets(buffer.getvalue())
    dataframe = read_excel_bytes(buffer.getvalue(), sheet_name="数据")

    assert sheets == ["说明", "数据"]
    assert dataframe.columns.tolist() == ["名称", "金额"]
    assert dataframe.iloc[0, 1] == "=SUM(B3:B4)"


def test_unknown_selected_column_has_clear_error() -> None:
    with pytest.raises(ValueError, match="找不到"):
        drop_duplicate_rows(pd.DataFrame({"A": [1]}), subset=["不存在"])
